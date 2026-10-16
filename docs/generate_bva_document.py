#!/usr/bin/env python3
"""Generate BVA document (Word + PDF) for TC-15 through TC-26."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUT_DIR = Path(__file__).resolve().parent
JSON_PATH = OUT_DIR / "bva_tc15_26_results.json"
DOCX_PATH = OUT_DIR / "Pengujian_Nilai_Batas_BVA.docx"
PDF_PATH = OUT_DIR / "Pengujian_Nilai_Batas_BVA.pdf"

NAVY = RGBColor(0x1B, 0x3A, 0x5F)
NAVY_HEX = colors.HexColor("#1B3A5F")
GOLD_HEX = colors.HexColor("#C9A227")
ROW_ALT = colors.HexColor("#F4F7FB")
PASS_BG = colors.HexColor("#E7F6EC")
FAIL_BG = colors.HexColor("#FDECEC")
HEADER_BG = colors.HexColor("#1B3A5F")

APP_NAME = "Aplikasi Simpan Pinjam Nasabah PT Puninar Jaya"
ORG_NAME = "Koperasi Konsumen Karyawan PT. Puninar Jaya (KOPPUJA)"
TODAY = "17 Agustus 2026"

HEADERS = [
    "ID Uji",
    "Skenario pengujian nilai batas",
    "Hasil yang diharapkan",
    "Hasil aktual",
    "Lulus / Gagal",
]


def load_cases():
    data = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    return data


def set_run_font(run, size=12, bold=False, color=None, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_p(doc, text, *, size=12, bold=False, align="justify", space_after=8, space_before=0, italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, color=NAVY)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    return p


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "1B3A5F")
        borders.append(el)
    tcPr.append(borders)


def write_cell(cell, text, *, bold=False, size=9, fill=None, center=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(str(text))
    text_color = RGBColor(255, 255, 255) if fill == "1B3A5F" else (color or RGBColor(0x22, 0x22, 0x22))
    set_run_font(run, size=size, bold=bold, color=text_color)
    if fill:
        shade_cell(cell, fill)
    set_cell_border(cell)


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Pengujian Nilai Batas TC-15 s.d. TC-26  |  " + APP_NAME + "  |  Halaman ")
    set_run_font(run, size=8, color=NAVY)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run2 = p.add_run()
    run2._r.append(fld1)
    run2._r.append(instr)
    run2._r.append(fld2)
    set_run_font(run2, size=8, color=NAVY)


def table_rows(cases):
    rows = []
    for case in cases:
        rows.append([
            case["id"],
            case["scenario"],
            case["expected"],
            case["actual"],
            case["status"],
        ])
    return rows


def build_docx(payload):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    add_page_number(section)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    add_p(doc, "DOKUMEN PENGUJIAN PERANGKAT LUNAK", size=12, bold=True, align="center", space_after=4, color=NAVY)
    add_p(doc, "PENGUJIAN NILAI BATAS", size=22, bold=True, align="center", space_after=2, color=NAVY)
    add_p(doc, "Boundary Value Analysis (BVA) — TC-15 s.d. TC-26", size=14, bold=True, align="center", space_after=12, color=RGBColor(0xC9, 0xA2, 0x27))
    add_p(doc, APP_NAME, size=13, bold=True, align="center", space_after=2)
    add_p(doc, ORG_NAME, size=12, italic=True, align="center", space_after=14)

    meta = doc.add_table(rows=5, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_rows = [
        ("Jenis pengujian", "Black-box — Boundary Value Analysis"),
        ("Kasus uji", "TC-15 sampai TC-26 (12 kasus)"),
        ("Objek uji", "RuleEngine (R01–R16) dan validasi formulir apply-loan.php"),
        ("Pelaksanaan", payload["executed_at"] + f" — {payload['passed']}/{payload['total']} lulus"),
        ("Versi dokumen", "2.0"),
    ]
    for i, (k, v) in enumerate(meta_rows):
        write_cell(meta.rows[i].cells[0], k, bold=True, size=10, fill="1B3A5F", center=True)
        write_cell(meta.rows[i].cells[1], v, size=10, fill="F4F7FB")

    heading(doc, "1. Pendahuluan")
    add_p(
        doc,
        "Pengujian nilai batas memeriksa perilaku sistem tepat pada ambang kebijakan pinjaman KOPPUJA. "
        "Setiap kasus membandingkan nilai sah pada batas (n) dengan nilai tepat di luar batas (n−1 atau n+1). "
        "Kasus TC-15 sampai TC-26 mencakup kelayakan pengajuan, keputusan persetujuan otomatis, dan validasi formulir.",
    )

    heading(doc, "2. Objek dan nilai batas yang diuji")
    ident_data = [
        ["TC-15", "R01", "masa_kerja_bulan", "12 bulan", "12 dan 11"],
        ["TC-16", "R08", "simpanan_wajib", "Rp 50.000", "50.000 dan 49.999"],
        ["TC-17", "R05 (2×)", "jumlah_pinjaman", "Rp 4.000.000", "4.000.000 dan 4.000.001"],
        ["TC-18", "R05 (1×)", "jumlah_pinjaman", "Rp 2.000.000", "2.000.000 dan 2.000.001"],
        ["TC-19", "R06", "tenor_bulan", "18 bulan", "18 dan 19"],
        ["TC-20", "R07", "rasio_angsuran_gaji", "40%", "40 dan 41"],
        ["TC-21", "R11", "jumlah_pinjaman", "Rp 5.000.000", "5.000.000 dan 5.000.001"],
        ["TC-22", "R12 / R15", "tenor_bulan", "12 bulan", "12 dan 13"],
        ["TC-23", "R13", "riwayat_tunggakan", "3 kali", "2 dan 3"],
        ["TC-24", "Formulir", "loan_amount", "Rp 100.000", "100.000 dan 99.999"],
        ["TC-25", "Formulir", "term_months", "1 dan 18", "0, 1, 18, 19"],
        ["TC-26", "R10 / R16", "auto_approve / potongan_hrd", "1 dan 0", "1 dan 0"],
    ]
    ident = doc.add_table(rows=1 + len(ident_data), cols=5)
    ident.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["ID", "Aturan", "Variabel", "Batas n", "Nilai uji"]):
        write_cell(ident.rows[0].cells[i], h, bold=True, size=9, fill="1B3A5F", center=True)
    for r, row in enumerate(ident_data):
        fill = "F4F7FB" if r % 2 == 0 else "FFFFFF"
        for c, val in enumerate(row):
            write_cell(ident.rows[r + 1].cells[c], val, size=9, fill=fill, center=c != 2)

    heading(doc, "3. Hasil pelaksanaan TC-15 s.d. TC-26")
    add_p(
        doc,
        f"Pengujian dijalankan terhadap RuleEngine dan fungsi validateLoanFormInput() pada {payload['executed_at']}. "
        f"Hasil: {payload['passed']} dari {payload['total']} kasus lulus.",
    )

    rows = table_rows(payload["cases"])
    table = doc.add_table(rows=1 + len(rows), cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [1.8, 5.2, 4.4, 5.0, 2.0]
    for i, h in enumerate(HEADERS):
        write_cell(table.rows[0].cells[i], h, bold=True, size=9, fill="1B3A5F", center=True)
        table.rows[0].cells[i].width = Cm(widths[i])
    for r, row in enumerate(rows):
        fill = "E7F6EC" if row[4] == "Lulus" else "FDECEC"
        for c, val in enumerate(row):
            write_cell(table.rows[r + 1].cells[c], val, size=8, fill=fill, center=(c in (0, 4)), bold=(c == 4))
            table.rows[r + 1].cells[c].width = Cm(widths[c])

    heading(doc, "4. Ringkasan")
    add_p(
        doc,
        "Seluruh titik batas kebijakan terbukti berperilaku sesuai rancangan: nilai tepat pada batas yang memakai operator “lebih dari” (>) masih diterima, "
        "sedangkan kelebihan satu satuan ditolak atau diteruskan ke peninjauan manual. Operator “lebih dari atau sama dengan” (>=) pada R13 memicu review mulai nilai 3. "
        "Validasi formulir menolak nominal di bawah Rp 100.000 serta tenor di luar 1–18 bulan.",
    )
    add_p(doc, f"Status akhir: {payload['passed']}/{payload['total']} kasus LULUS.", bold=True, align="center", space_before=10)
    add_p(doc, TODAY, align="center", size=11)

    doc.save(DOCX_PATH)


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoverKicker", fontName="Times-Bold", fontSize=11, leading=14, alignment=TA_CENTER, textColor=NAVY_HEX, spaceAfter=6))
    styles.add(ParagraphStyle(name="CoverTitle", fontName="Times-Bold", fontSize=20, leading=24, alignment=TA_CENTER, textColor=NAVY_HEX, spaceAfter=4))
    styles.add(ParagraphStyle(name="CoverSub", fontName="Times-Bold", fontSize=12, leading=16, alignment=TA_CENTER, textColor=GOLD_HEX, spaceAfter=12))
    styles.add(ParagraphStyle(name="CoverApp", fontName="Times-Bold", fontSize=12, leading=16, alignment=TA_CENTER, spaceAfter=2))
    styles.add(ParagraphStyle(name="CoverOrg", fontName="Times-Italic", fontSize=11, leading=14, alignment=TA_CENTER, textColor=colors.HexColor("#333333"), spaceAfter=14))
    styles.add(ParagraphStyle(name="BodyJust", fontName="Times-Roman", fontSize=11, leading=15, alignment=TA_JUSTIFY, spaceAfter=8))
    styles.add(ParagraphStyle(name="H1", fontName="Times-Bold", fontSize=14, leading=18, textColor=NAVY_HEX, spaceBefore=12, spaceAfter=6))
    styles.add(ParagraphStyle(name="CellH", fontName="Times-Bold", fontSize=8, leading=10, textColor=colors.white, alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="Cell", fontName="Times-Roman", fontSize=8, leading=11, alignment=TA_LEFT, textColor=colors.HexColor("#222222")))
    styles.add(ParagraphStyle(name="CellC", fontName="Times-Bold", fontSize=8, leading=11, alignment=TA_CENTER, textColor=colors.HexColor("#1B5E20")))
    styles.add(ParagraphStyle(name="BulletBody", fontName="Times-Roman", fontSize=11, leading=15, alignment=TA_JUSTIFY))
    return styles


def make_pdf_table(headers, rows, col_widths, status_col=None):
    styles = pdf_styles()
    data = [headers] + rows
    formatted = [[Paragraph(str(c), styles["CellH"]) for c in data[0]]]
    for row in data[1:]:
        cells = []
        for i, val in enumerate(row):
            st = styles["CellC"] if status_col is not None and i == status_col else styles["Cell"]
            cells.append(Paragraph(str(val), st))
        formatted.append(cells)
    t = Table(formatted, colWidths=col_widths, repeatRows=1)
    cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), HEADER_BG),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.3, NAVY_HEX),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
    ]
    for i, row in enumerate(data[1:], start=1):
        bg = PASS_BG if (status_col is not None and row[status_col] == "Lulus") else (FAIL_BG if status_col is not None else (ROW_ALT if i % 2 else colors.white))
        if status_col is None:
            bg = ROW_ALT if i % 2 else colors.white
        cmds.append(("BACKGROUND", (0, i), (-1, i), bg))
    t.setStyle(TableStyle(cmds))
    return t


def on_page(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(NAVY_HEX)
    canvas.setLineWidth(1.2)
    canvas.line(1.8 * cm, A4[1] - 1.4 * cm, A4[0] - 1.8 * cm, A4[1] - 1.4 * cm)
    canvas.setFont("Times-Italic", 8)
    canvas.setFillColor(NAVY_HEX)
    canvas.drawString(1.8 * cm, A4[1] - 1.25 * cm, "Pengujian Nilai Batas TC-15 s.d. TC-26")
    canvas.drawRightString(A4[0] - 1.8 * cm, A4[1] - 1.25 * cm, "KOPPUJA")
    canvas.line(1.8 * cm, 1.4 * cm, A4[0] - 1.8 * cm, 1.4 * cm)
    canvas.drawCentredString(A4[0] / 2, 0.85 * cm, f"{APP_NAME}  |  Halaman {doc.page}")
    canvas.restoreState()


def xml_escape(text):
    return (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def build_pdf(payload):
    styles = pdf_styles()
    story = []
    usable = A4[0] - 3.6 * cm

    story.append(Paragraph("DOKUMEN PENGUJIAN PERANGKAT LUNAK", styles["CoverKicker"]))
    story.append(Paragraph("PENGUJIAN NILAI BATAS", styles["CoverTitle"]))
    story.append(Paragraph("Boundary Value Analysis (BVA) — TC-15 s.d. TC-26", styles["CoverSub"]))
    story.append(Paragraph(APP_NAME, styles["CoverApp"]))
    story.append(Paragraph(ORG_NAME, styles["CoverOrg"]))

    story.append(make_pdf_table(
        ["Atribut", "Keterangan"],
        [
            ["Jenis pengujian", "Black-box — Boundary Value Analysis"],
            ["Kasus uji", "TC-15 sampai TC-26 (12 kasus)"],
            ["Objek uji", "RuleEngine (R01–R16) dan validasi formulir apply-loan.php"],
            ["Pelaksanaan", f"{payload['executed_at']} — {payload['passed']}/{payload['total']} lulus"],
            ["Versi dokumen", "2.0"],
        ],
        [4.5 * cm, usable - 4.5 * cm],
    ))

    story.append(Paragraph("1. Pendahuluan", styles["H1"]))
    story.append(Paragraph(
        "Pengujian nilai batas memeriksa perilaku sistem tepat pada ambang kebijakan pinjaman KOPPUJA. "
        "Setiap kasus membandingkan nilai sah pada batas (n) dengan nilai tepat di luar batas (n−1 atau n+1). "
        "Kasus TC-15 sampai TC-26 mencakup kelayakan pengajuan, keputusan persetujuan otomatis, dan validasi formulir.",
        styles["BodyJust"],
    ))

    story.append(Paragraph("2. Objek dan nilai batas yang diuji", styles["H1"]))
    story.append(make_pdf_table(
        ["ID", "Aturan", "Variabel", "Batas n", "Nilai uji"],
        [
            ["TC-15", "R01", "masa_kerja_bulan", "12 bulan", "12 dan 11"],
            ["TC-16", "R08", "simpanan_wajib", "Rp 50.000", "50.000 dan 49.999"],
            ["TC-17", "R05 (2×)", "jumlah_pinjaman", "Rp 4.000.000", "4.000.000 dan 4.000.001"],
            ["TC-18", "R05 (1×)", "jumlah_pinjaman", "Rp 2.000.000", "2.000.000 dan 2.000.001"],
            ["TC-19", "R06", "tenor_bulan", "18 bulan", "18 dan 19"],
            ["TC-20", "R07", "rasio_angsuran_gaji", "40%", "40 dan 41"],
            ["TC-21", "R11", "jumlah_pinjaman", "Rp 5.000.000", "5.000.000 dan 5.000.001"],
            ["TC-22", "R12 / R15", "tenor_bulan", "12 bulan", "12 dan 13"],
            ["TC-23", "R13", "riwayat_tunggakan", "3 kali", "2 dan 3"],
            ["TC-24", "Formulir", "loan_amount", "Rp 100.000", "100.000 dan 99.999"],
            ["TC-25", "Formulir", "term_months", "1 dan 18", "0, 1, 18, 19"],
            ["TC-26", "R10 / R16", "auto_approve / potongan_hrd", "1 dan 0", "1 dan 0"],
        ],
        [1.8 * cm, 2.4 * cm, 4.2 * cm, 3.2 * cm, usable - 11.6 * cm],
    ))

    story.append(Paragraph("3. Hasil pelaksanaan TC-15 s.d. TC-26", styles["H1"]))
    story.append(Paragraph(
        f"Pengujian dijalankan terhadap RuleEngine dan fungsi validateLoanFormInput() pada {xml_escape(payload['executed_at'])}. "
        f"Hasil: {payload['passed']} dari {payload['total']} kasus lulus.",
        styles["BodyJust"],
    ))

    result_rows = [
        [xml_escape(c["id"]), xml_escape(c["scenario"]), xml_escape(c["expected"]), xml_escape(c["actual"]), xml_escape(c["status"])]
        for c in payload["cases"]
    ]
    story.append(make_pdf_table(
        HEADERS,
        result_rows,
        [1.6 * cm, 4.4 * cm, 4.0 * cm, 5.2 * cm, 1.8 * cm],
        status_col=4,
    ))

    story.append(Paragraph("4. Ringkasan", styles["H1"]))
    story.append(Paragraph(
        "Seluruh titik batas kebijakan terbukti berperilaku sesuai rancangan: nilai tepat pada batas yang memakai operator "
        "“lebih dari” (&gt;) masih diterima, sedangkan kelebihan satu satuan ditolak atau diteruskan ke peninjauan manual. "
        "Operator “lebih dari atau sama dengan” (&gt;=) pada R13 memicu review mulai nilai 3. "
        "Validasi formulir menolak nominal di bawah Rp 100.000 serta tenor di luar 1–18 bulan.",
        styles["BodyJust"],
    ))
    story.append(Spacer(1, 8))
    story.append(Paragraph(f"Status akhir: {payload['passed']}/{payload['total']} kasus LULUS.", styles["CoverApp"]))
    story.append(Paragraph(TODAY, styles["CoverOrg"]))

    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        leftMargin=1.8 * cm,
        rightMargin=1.8 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
        title="Pengujian Nilai Batas TC-15 s.d. TC-26",
        author=ORG_NAME,
        subject="Boundary Value Analysis",
    )
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


if __name__ == "__main__":
    payload = load_cases()
    build_docx(payload)
    build_pdf(payload)
    print(f"Wrote {DOCX_PATH}")
    print(f"Wrote {PDF_PATH}")
